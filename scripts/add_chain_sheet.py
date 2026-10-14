"""Add chain_economics.csv to the workbook as sheet 'chain_economics' (run after build_merged.py)."""
import sys
import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
import pandas as pd

WB, CSV, LEVEL = sys.argv[1], sys.argv[2], sys.argv[3]
df = pd.read_csv(CSV)
wb = openpyxl.load_workbook(WB)
if "chain_economics" in wb.sheetnames:
    del wb["chain_economics"]
ws = wb.create_sheet("chain_economics", 1)
ws.append(list(df.columns))
for c in ws[1]:
    c.font = Font(bold=True)
for row in df.itertuples(index=False):
    ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in row])
ws.freeze_panes = "B2"
ws.auto_filter.ref = ws.dimensions
for i, c in enumerate(df.columns, 1):
    ws.column_dimensions[get_column_letter(i)].width = 30 if c in ("Product", "Ingredients", "Check vs hayday.info") else 14

n = wb["Notes"]
n.append([])
n.append(["Chain economics (chain_economics) - planner step 1"]); n.cell(n.max_row, 1).font = Font(bold=True)
for k, v in [
    ("What", "Per 1 unit of each product, over its whole ingredient chain: resource-minutes by type, total chain minutes, critical path, ingredient value, coins/hr. Script: scripts/chain_economics.py."),
    ("Price", "verified_in_game.csv first, then hayday.info, then Fandom max price. Apple Juice = 129 confirmed in game 2026-10-14."),
    ("Times", "Fandom normal times; '... mastered' columns use Fandom mastered time for machine steps only (crops/animals have no mastery)."),
    ("Crops", "1 field-cycle = net 1 crop (plant 1, harvest 2). Field-min per crop = grow time."),
    ("Trees", "Assumption: 3 fruit per harvest (wiki: harvests give 2, 3, 4). Tree-min per fruit = time / 3. Tree purchase cost ignored."),
    ("Animals", "1 good per feed per animal timer. Feed links (Egg<-Chicken Feed, Milk<-Cow Feed, Bacon<-Pig Feed, Wool<-Sheep Feed, Goat Milk<-Goat Feed, Lamb Chop<-Lamb Feed) typed from Fandom page text. Honeycomb = beehive timer, no feed."),
    ("Feed batches", "Feed Mill makes 3 sacks per batch (Fandom 'three by three'); recipe quantities and time divided by 3 per sack."),
    ("External", "Mine ores, fish, lobster, duck feather, mussel, peanuts need tools/lures/nets: counted at market value with 0 production time. Products using them look better than they are (column 'External inputs')."),
    ("Diamond", "Diamond Ring and Gracious Bouquet use diamonds (premium currency) - not costed."),
    ("Critical path", "Wall-clock from empty barn: unlimited fields/animals in parallel, one machine of each kind, ingredient batches back to back. Ignores your check-in interval (planner step 2)."),
    ("Coins/hr per chain-min", "Price / (machine + field + animal + tree minutes) x 60. Weighs a field-minute the same as a machine-minute - a rough ranking only; the LP planner weighs them by what you actually have."),
    ("Value added/hr own machine", "(Price - ingredient value) / own machine time x 60. Same formula as hayday.info profit_per_hour: 273 of 289 rows match (site uses unrounded mastered time = normal x 0.85, or normal time). 16 rows differ - see 'Check vs hayday.info'."),
    ("Unlocked", f"'Unlocked at L{LEVEL}' = product and every ingredient in its chain have Fandom level <= {LEVEL}. Unknown level = not unlocked."),
]:
    n.append([k, v])
wb.save(WB)
print("sheet added:", len(df), "rows")
